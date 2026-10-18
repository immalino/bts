"""
Unit tests untuk src/utils/excel_utils.py dan src/core/excel_generator.py (Sesi 9).

Memverifikasi:
1. Kalkulasi dimensi sel tunggal dan merged cell di openpyxl.
2. Kalkulasi proporsional fitted dimensions dengan aspect ratio dan padding.
3. Pembuatan file Excel laporan berbasis template multi-sheet dengan gambar tertanam.
4. Dukungan EXIF orientation transpose pada gambar.
5. Mode ukuran fit_to_cell dan custom sizing.
6. Penanganan error (template hilang, sheet hilang, foto tidak ditemukan/korup).
7. Pembatalan (cancellation) proses pembuatan dokumen.
"""

import os
from pathlib import Path
import openpyxl
from PIL import Image
import pytest

from src.core.excel_generator import ExcelGeneratorThread
from src.core.models import PhotoFieldConfig, ProjectData, TemplateConfig
from src.utils.excel_utils import (
    calculate_fitted_dimensions,
    calculate_image_anchor,
    format_output_filename,
    get_cell_dimensions,
    sanitize_excel_output,
)


# =====================================================================
# Fixtures
# =====================================================================

@pytest.fixture
def dummy_template_path(tmp_path: Path) -> str:
    """Membuat file template Excel dummy dengan 3 sheet dan sel merged."""
    wb = openpyxl.Workbook()
    # Sheet 1: Antenna
    ws1 = wb.active
    ws1.title = "Antenna"
    ws1["A1"] = "LAPORAN FOTO ANTENA"
    ws1.column_dimensions["B"].width = 20.0
    ws1.column_dimensions["C"].width = 20.0
    ws1.column_dimensions["D"].width = 20.0
    ws1.column_dimensions["E"].width = 20.0
    for r in range(5, 16):
        ws1.row_dimensions[r].height = 25.0
    ws1.merge_cells("B5:E15")

    # Sheet 2: Equipment
    ws2 = wb.create_sheet(title="Equipment")
    ws2["A1"] = "FOTO PERANGKAT DAN RACK"
    ws2.column_dimensions["B"].width = 30.0
    ws2.row_dimensions[4].height = 50.0

    # Sheet 3: Tower
    ws3 = wb.create_sheet(title="Tower")
    ws3["A1"] = "FOTO STRUKTUR TOWER"
    ws3.column_dimensions["C"].width = 25.0
    ws3.row_dimensions[6].height = 40.0

    template_file = str(tmp_path / "dummy_template.xlsx")
    wb.save(template_file)
    return template_file


@pytest.fixture
def sample_photos(tmp_path: Path) -> dict:
    """Membuat beberapa file foto dummy (standar, portrait, EXIF orientation)."""
    photos_dir = tmp_path / "photos"
    photos_dir.mkdir()

    # 1. Foto standar landscape 800x600
    p1 = photos_dir / "antenna_1.jpg"
    img1 = Image.new("RGB", (800, 600), color="blue")
    img1.save(str(p1), format="JPEG")

    # 2. Foto portrait 600x800
    p2 = photos_dir / "equipment_1.jpg"
    img2 = Image.new("RGB", (600, 800), color="green")
    img2.save(str(p2), format="JPEG")

    # 3. Foto dengan tag EXIF orientation (6 = Rotate 90 CW)
    p3 = photos_dir / "tower_1.jpg"
    img3 = Image.new("RGB", (1000, 500), color="red")
    exif = img3.getexif()
    exif[0x0112] = 6
    img3.save(str(p3), format="JPEG", exif=exif)

    return {
        "p1": str(p1),
        "p2": str(p2),
        "p3": str(p3),
    }


# =====================================================================
# Test Excel Utils
# =====================================================================

class TestExcelUtils:
    def test_single_cell_dimensions(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.column_dimensions["B"].width = 20.0
        ws.row_dimensions[5].height = 30.0

        w_px, h_px = get_cell_dimensions(ws, "B5")
        # lebar: 20 * 7 + 5 = 145 px
        # tinggi: round(30 * 96 / 72) = 40 px
        assert w_px == 145
        assert h_px == 40

    def test_merged_cell_dimensions(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        # Merge B5:D6 (3 kolom: B, C, D; 2 baris: 5, 6)
        ws.column_dimensions["B"].width = 10.0  # 10*7+5 = 75 px
        ws.column_dimensions["C"].width = 10.0  # 75 px
        ws.column_dimensions["D"].width = 10.0  # 75 px
        ws.row_dimensions[5].height = 20.0     # round(20 * 4/3) = 27 px
        ws.row_dimensions[6].height = 20.0     # 27 px

        ws.merge_cells("B5:D6")

        # Cek koordinat awal merged cell
        w_px, h_px = get_cell_dimensions(ws, "B5")
        assert w_px == 75 * 3  # 225 px
        assert h_px == 27 * 2  # 54 px

        # Cek sel di dalam rentang merged cell juga menghasilkan dimensi total
        w_inner, h_inner = get_cell_dimensions(ws, "C6")
        assert w_inner == 225
        assert h_inner == 54

    def test_coordinate_normalization_and_invalid(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.column_dimensions["A"].width = 15.0
        ws.row_dimensions[1].height = 15.0

        # Case-insensitive dan whitespace
        w1, h1 = get_cell_dimensions(ws, " a1 ")
        w2, h2 = get_cell_dimensions(ws, "$A$1")
        assert w1 == w2
        assert h1 == h2

        # Invalid coordinate
        with pytest.raises(ValueError):
            get_cell_dimensions(ws, "")

        with pytest.raises(ValueError):
            get_cell_dimensions(ws, "INVALID_CELL")

    def test_calculate_fitted_dimensions_landscape(self):
        # Gambar landscape 800x400 (rasio 2:1), sel 200x200, padding 2
        # Ruang tersedia: 200 - 4 = 196
        # Skala: min(196/800, 196/400) = 196/800 = 0.245
        # new_w = 196, new_h = round(400 * 0.245) = 98
        w, h = calculate_fitted_dimensions(800, 400, 200, 200, padding=2)
        assert w == 196
        assert h == 98
        assert w <= 196
        assert h <= 196

    def test_calculate_fitted_dimensions_portrait(self):
        # Gambar portrait 400x800 (rasio 1:2), sel 200x200, padding 0
        w, h = calculate_fitted_dimensions(400, 800, 200, 200, padding=0)
        assert w == 100
        assert h == 200

    def test_calculate_fitted_dimensions_edge_cases(self):
        # Ukuran nol / negatif
        w, h = calculate_fitted_dimensions(0, 0, 100, 100, padding=2)
        assert w == 96
        assert h == 96

        # Padding melebihi ukuran kotak
        w, h = calculate_fitted_dimensions(500, 500, 10, 10, padding=20)
        assert w >= 1
        assert h >= 1

    def test_calculate_image_anchor_single_cell_center(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.column_dimensions["B"].width = 20.0  # round(20*7+5) = 145 px
        ws.row_dimensions[5].height = 50.0      # round(50*96/72) = 67 px

        anchor = calculate_image_anchor(ws, "B5", 45, 27, alignment="center")
        # Col B is index 1, row 5 is index 4
        assert anchor._from.col == 1
        assert anchor._from.row == 4
        # Expected offsets: (145 - 45) // 2 = 50 px; (67 - 27) // 2 = 20 px
        assert anchor._from.colOff == 50 * 9525
        assert anchor._from.rowOff == 20 * 9525
        assert anchor.ext.cx == 45 * 9525
        assert anchor.ext.cy == 27 * 9525

    def test_calculate_image_anchor_single_cell_top_left(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.column_dimensions["B"].width = 20.0
        ws.row_dimensions[5].height = 50.0

        anchor = calculate_image_anchor(ws, "B5", 45, 27, alignment="top-left")
        assert anchor._from.col == 1
        assert anchor._from.row == 4
        assert anchor._from.colOff == 0
        assert anchor._from.rowOff == 0

    def test_calculate_image_anchor_merged_cell_center(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.column_dimensions["B"].width = 20.0  # 145 px
        ws.column_dimensions["C"].width = 20.0  # 145 px -> total 290 px
        ws.row_dimensions[5].height = 30.0      # 40 px
        ws.row_dimensions[6].height = 30.0      # 40 px -> total 80 px
        ws.merge_cells("B5:C6")

        # Gambar 90 x 40 px: offset_x = (290 - 90)//2 = 100 px (within B: 100 < 145)
        # offset_y = (80 - 40)//2 = 20 px (within Row 5: 20 < 40)
        anchor = calculate_image_anchor(ws, "B5", 90, 40, alignment="center")
        assert anchor._from.col == 1
        assert anchor._from.row == 4
        assert anchor._from.colOff == 100 * 9525
        assert anchor._from.rowOff == 20 * 9525

    def test_calculate_image_anchor_invalid(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        with pytest.raises(ValueError):
            calculate_image_anchor(ws, "", 100, 100)
        with pytest.raises(ValueError):
            calculate_image_anchor(ws, "INVALID_XYZ", 100, 100)


# =====================================================================
# Test ExcelGeneratorThread
# =====================================================================

class TestExcelGeneratorThread:
    def test_generate_multi_sheet_with_merged_cells(
        self, tmp_path: Path, dummy_template_path: str, sample_photos: dict
    ):
        """
        Pengujian pembuatan file Excel dengan 3 sheet dan sel merged.
        Verifikasi gambar terpasang di sheet yang sesuai dan dimensi sel dihormati.
        """
        output_file = str(tmp_path / "output_report.xlsx")

        fields = [
            PhotoFieldConfig(
                field_id="f_antenna",
                label="Foto Antena",
                sheet_name="Antenna",
                cell_anchor="B5",
                sizing_mode="fit_to_cell",
                padding=2,
            ),
            PhotoFieldConfig(
                field_id="f_equip",
                label="Foto Equipment",
                sheet_name="Equipment",
                cell_anchor="B4",
                sizing_mode="fit_to_cell",
                padding=2,
            ),
            PhotoFieldConfig(
                field_id="f_tower",
                label="Foto Tower",
                sheet_name="Tower",
                cell_anchor="C6",
                sizing_mode="fit_to_cell",
                padding=2,
            ),
        ]

        photo_assignments = {
            "f_antenna": sample_photos["p1"],
            "f_equip": sample_photos["p2"],
            "f_tower": sample_photos["p3"],
        }

        template_config = TemplateConfig(
            template_id="tpl_multi",
            name="Template Multi Sheet",
            template_file_path=dummy_template_path,
            fields=fields,
        )

        progress_events = []
        finished_items = []
        completed_events = []

        thread = ExcelGeneratorThread(
            template=template_config,
            output_path=output_file,
            photo_assignments=photo_assignments,
        )
        thread.progress_updated.connect(lambda s, t, m: progress_events.append((s, t, m)))
        thread.item_finished.connect(lambda fid, s, m: finished_items.append((fid, s, m)))
        thread.generation_completed.connect(
            lambda out, sc, fc, errs: completed_events.append((out, sc, fc, errs))
        )

        # Eksekusi thread
        thread.run()

        # 1. Verifikasi completion signal
        assert len(completed_events) == 1
        out_p, success_cnt, fail_cnt, errors = completed_events[0]
        assert out_p == output_file
        assert success_cnt == 3
        assert fail_cnt == 0
        assert errors == []

        # 2. Verifikasi item_finished signals
        assert len(finished_items) == 3
        for fid, success, _ in finished_items:
            assert success is True
            assert fid in photo_assignments

        # 3. Verifikasi progress signals
        assert len(progress_events) >= 4

        # 4. Verifikasi file fisik terbentuk
        assert os.path.exists(output_file)
        assert os.path.getsize(output_file) > 0

        # 5. Reload workbook dan verifikasi gambar tertanam di tiap sheet
        wb_out = openpyxl.load_workbook(output_file)
        assert "Antenna" in wb_out.sheetnames
        assert "Equipment" in wb_out.sheetnames
        assert "Tower" in wb_out.sheetnames

        ws1 = wb_out["Antenna"]
        assert len(ws1._images) == 1
        img1 = ws1._images[0]
        # B5: col B (index 1), row 5 (index 4)
        assert img1.anchor._from.col == 1
        assert img1.anchor._from.row == 4

        ws2 = wb_out["Equipment"]
        assert len(ws2._images) == 1
        img2 = ws2._images[0]
        # B4: col B (index 1), row 4 (index 3)
        assert img2.anchor._from.col == 1
        assert img2.anchor._from.row == 3

        ws3 = wb_out["Tower"]
        assert len(ws3._images) == 1
        img3 = ws3._images[0]
        # C6: col C (index 2), row 6 (index 5)
        assert img3.anchor._from.col == 2
        assert img3.anchor._from.row == 5
        wb_out.close()

    def test_custom_sizing_mode(
        self, tmp_path: Path, dummy_template_path: str, sample_photos: dict
    ):
        """Verifikasi mode ukuran kustom (custom_width & custom_height)."""
        output_file = str(tmp_path / "output_custom.xlsx")

        field = PhotoFieldConfig(
            field_id="f_custom",
            label="Foto Custom Size",
            sheet_name="Equipment",
            cell_anchor="B2",
            sizing_mode="custom",
            custom_width=150,
            custom_height=120,
        )

        thread = ExcelGeneratorThread(
            template=dummy_template_path,
            output_path=output_file,
            fields=[field],
            photo_assignments={"f_custom": sample_photos["p1"]},
        )
        thread.run()

        wb = openpyxl.load_workbook(output_file)
        ws = wb["Equipment"]
        assert len(ws._images) == 1
        img = ws._images[0]
        # 1 px = 9525 EMUs
        assert round(img.anchor.ext.cx / 9525) == 150
        assert round(img.anchor.ext.cy / 9525) == 120
        wb.close()

    def test_with_project_data(
        self, tmp_path: Path, dummy_template_path: str, sample_photos: dict
    ):
        """Verifikasi inisialisasi menggunakan model ProjectData."""
        out_dir = tmp_path / "project_out"
        out_dir.mkdir()

        field = PhotoFieldConfig(
            field_id="sec1",
            label="Sector 1",
            sheet_name="Antenna",
            cell_anchor="B5",
        )
        template_cfg = TemplateConfig(
            template_id="t1",
            name="T1",
            template_file_path=dummy_template_path,
            fields=[field],
        )
        proj = ProjectData(
            project_id="p100",
            site_name="SITE_JKT_01",
            template_id="t1",
            photo_assignments={"sec1": sample_photos["p1"]},
            output_dir=str(out_dir),
        )

        thread = ExcelGeneratorThread(template=template_cfg, project_data=proj)
        assert thread.output_path == str(out_dir / "Laporan_SITE_JKT_01.xlsx")
        thread.run()

        assert os.path.exists(thread.output_path)

    def test_error_missing_template(self, tmp_path: Path):
        """Verifikasi penanganan error jika file template tidak ditemukan."""
        output_file = str(tmp_path / "out.xlsx")
        thread = ExcelGeneratorThread(
            template="/non_existent/template.xlsx",
            output_path=output_file,
            fields=[],
        )
        completed = []
        thread.generation_completed.connect(lambda o, s, f, err: completed.append((o, s, f, err)))
        thread.run()

        assert len(completed) == 1
        _, success, fail, errors = completed[0]
        assert success == 0
        assert len(errors) > 0
        assert "tidak ditemukan" in errors[0]

    def test_error_missing_sheet_and_missing_photo(
        self, tmp_path: Path, dummy_template_path: str
    ):
        """Verifikasi pelaporan kegagalan per field jika sheet atau foto hilang."""
        output_file = str(tmp_path / "out.xlsx")
        fields = [
            PhotoFieldConfig(
                field_id="f_no_sheet",
                label="No Sheet",
                sheet_name="NonExistentSheet",
                cell_anchor="A1",
            ),
            PhotoFieldConfig(
                field_id="f_no_photo",
                label="No Photo",
                sheet_name="Antenna",
                cell_anchor="B5",
            ),
            PhotoFieldConfig(
                field_id="f_invalid_file",
                label="Invalid File",
                sheet_name="Antenna",
                cell_anchor="B5",
            ),
        ]
        photo_assignments = {
            "f_no_sheet": "dummy.jpg",
            # f_no_photo unassigned
            "f_invalid_file": "/path/to/missing_file.jpg",
        }

        finished = []
        completed = []
        thread = ExcelGeneratorThread(
            template=dummy_template_path,
            output_path=output_file,
            fields=fields,
            photo_assignments=photo_assignments,
        )
        thread.item_finished.connect(lambda fid, s, m: finished.append((fid, s, m)))
        thread.generation_completed.connect(lambda o, s, f, err: completed.append((o, s, f, err)))
        thread.run()

        assert len(finished) == 3
        for _, s, _ in finished:
            assert s is False

        _, success, fail, errors = completed[0]
        assert success == 0
        assert fail == 3
        assert len(errors) == 3

    def test_error_corrupt_photo_file(
        self, tmp_path: Path, dummy_template_path: str
    ):
        """Verifikasi penanganan gambar yang korup."""
        output_file = str(tmp_path / "out.xlsx")
        corrupt_img = tmp_path / "corrupt.jpg"
        corrupt_img.write_bytes(b"CORRUPT_NOT_AN_IMAGE_DATA")

        field = PhotoFieldConfig(
            field_id="corrupt_field",
            label="Corrupt Photo",
            sheet_name="Antenna",
            cell_anchor="B5",
        )

        finished = []
        thread = ExcelGeneratorThread(
            template=dummy_template_path,
            output_path=output_file,
            fields=[field],
            photo_assignments={"corrupt_field": str(corrupt_img)},
        )
        thread.item_finished.connect(lambda fid, s, m: finished.append((fid, s, m)))
        thread.run()

        assert len(finished) == 1
        assert finished[0][1] is False
        assert "Gagal memproses" in finished[0][2] or "Format foto tidak valid" in finished[0][2]

    def test_cancellation(
        self, tmp_path: Path, dummy_template_path: str, sample_photos: dict
    ):
        """Verifikasi fitur cancel() membatalkan proses generator."""
        output_file = str(tmp_path / "cancelled.xlsx")
        field = PhotoFieldConfig(
            field_id="f1",
            label="Field 1",
            sheet_name="Antenna",
            cell_anchor="B5",
        )
        thread = ExcelGeneratorThread(
            template=dummy_template_path,
            output_path=output_file,
            fields=[field],
            photo_assignments={"f1": sample_photos["p1"]},
        )
        thread.cancel()
        assert thread.is_cancelled is True

        completed = []
        thread.generation_completed.connect(lambda o, s, f, err: completed.append((o, s, f, err)))
        thread.run()

        assert len(completed) == 1
        _, success, fail, errors = completed[0]
        assert "dibatalkan" in errors[0]
        assert not os.path.exists(output_file)

    def test_high_resolution_image_preservation(
        self, tmp_path: Path, dummy_template_path: str
    ):
        """
        Memastikan foto beresolusi tinggi (misal 1600x1200) yang disematkan ke sel kecil
        tetap mempertahankan resolusi piksel fisiknya di dalam arsip Excel (tidak di-downsample
        hingga seukuran sel), sehingga foto tidak blur ketika dilihat atau di-zoom.
        """
        import io
        import zipfile

        output_file = str(tmp_path / "high_res_test.xlsx")
        high_res_photo = str(tmp_path / "large_photo.jpg")

        # Buat gambar resolusi 1600 x 1200
        Image.new("RGB", (1600, 1200), color="blue").save(
            high_res_photo, format="JPEG", quality=95
        )

        field = PhotoFieldConfig(
            field_id="f_highres",
            label="Foto High Res",
            sheet_name="Equipment",
            cell_anchor="B4",
            sizing_mode="fit_to_cell",
        )

        thread = ExcelGeneratorThread(
            template=dummy_template_path,
            output_path=output_file,
            fields=[field],
            photo_assignments={"f_highres": high_res_photo},
        )
        thread.run()

        assert os.path.exists(output_file)

        # 1. Periksa ukuran tampilan pada lembar kerja Excel (tetap fit di sel)
        wb = openpyxl.load_workbook(output_file)
        ws = wb["Equipment"]
        assert len(ws._images) == 1
        img = ws._images[0]
        disp_w = round(img.anchor.ext.cx / 9525)
        disp_h = round(img.anchor.ext.cy / 9525)
        assert disp_w <= 300
        assert disp_h <= 100
        wb.close()

        # 2. Buka arsip zip Excel dan verifikasi piksel fisik gambar media tetap 1600x1200 (tajam)
        with zipfile.ZipFile(output_file, "r") as z:
            media_files = [n for n in z.namelist() if n.startswith("xl/media/")]
            assert len(media_files) >= 1
            raw_bytes = z.read(media_files[0])
            with Image.open(io.BytesIO(raw_bytes)) as embedded_img:
                assert embedded_img.size == (1600, 1200)

    def test_excel_generator_with_text_fields(
        self,
        dummy_template_path: str,
        sample_photos: dict,
        tmp_path: Path,
    ):
        """Memverifikasi penulisan nilai field bertipe teks ke dalam sel Excel target."""
        output_file = str(tmp_path / "output_with_text.xlsx")

        field_photo = PhotoFieldConfig(
            field_id="f_ant",
            label="Foto Antena",
            sheet_name="Antenna",
            cell_anchor="B5",
            field_type="photo",
        )
        field_text1 = PhotoFieldConfig(
            field_id="t_site_name",
            label="Nama Site",
            sheet_name="Equipment",
            cell_anchor="A1",
            field_type="text",
        )
        field_text2 = PhotoFieldConfig(
            field_id="t_def",
            label="Catatan Default",
            sheet_name="Tower",
            cell_anchor="A1",
            field_type="text",
            default_text="Default Tower Note",
        )

        completed_data = {}

        def on_completed(path, success, fail, errors):
            completed_data["path"] = path
            completed_data["success"] = success
            completed_data["fail"] = fail
            completed_data["errors"] = errors

        thread = ExcelGeneratorThread(
            template=dummy_template_path,
            output_path=output_file,
            fields=[field_photo, field_text1, field_text2],
            photo_assignments={"f_ant": sample_photos["p1"]},
            text_assignments={"t_site_name": "JKT-KUNINGAN-001"},
        )
        thread.generation_completed.connect(on_completed)
        thread.run()

        assert os.path.exists(output_file)
        assert completed_data["success"] == 3
        assert completed_data["fail"] == 0

        # Verifikasi hasil isi file Excel
        wb = openpyxl.load_workbook(output_file)
        assert wb["Equipment"]["A1"].value == "JKT-KUNINGAN-001"
        assert wb["Tower"]["A1"].value == "Default Tower Note"
        assert len(wb["Antenna"]._images) == 1
        wb.close()

    def test_sanitize_excel_output_removes_calcchain_and_empty_v_tags(self, tmp_path: Path):
        """Memverifikasi sanitize_excel_output menghapus calcChain.xml dan memperbaiki empty <v />."""
        import zipfile
        from openpyxl.styles import Border, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = 100
        ws["A2"] = 200
        ws["A3"] = "=SUM(A1:A2)"
        thin = Side(border_style="thin", color="000000")
        ws["A3"].border = Border(top=thin, bottom=thin)

        raw_file = str(tmp_path / "raw_with_calcchain.xlsx")
        wb.save(raw_file)

        # Sisipkan tiruan calcChain.xml ke arsip zip untuk simulasi file buatan Microsoft Excel
        mod_file = str(tmp_path / "simulated_excel.xlsx")
        with zipfile.ZipFile(raw_file, "r") as zin, zipfile.ZipFile(mod_file, "w") as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "xl/_rels/workbook.xml.rels":
                    data = data.replace(
                        b"</Relationships>",
                        b'<Relationship Id="rIdCalc" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain" Target="calcChain.xml"/></Relationships>',
                    )
                elif item.filename == "[Content_Types].xml":
                    data = data.replace(
                        b"</Types>",
                        b'<Override PartName="/xl/calcChain.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/></Types>',
                    )
                zout.writestr(item, data)
            zout.writestr("xl/calcChain.xml", b"<calcChain><c r='A3' i='1'/></calcChain>")

        # Jalankan sanitasi
        sanitize_excel_output(mod_file)

        # Periksa arsip zip setelah sanitasi
        with zipfile.ZipFile(mod_file, "r") as zcheck:
            namelist = zcheck.namelist()
            assert "xl/calcChain.xml" not in namelist

            rels_text = zcheck.read("xl/_rels/workbook.xml.rels").decode("utf-8")
            assert "calcChain.xml" not in rels_text

            types_text = zcheck.read("[Content_Types].xml").decode("utf-8")
            assert "calcChain.xml" not in types_text

            sheet_text = zcheck.read("xl/worksheets/sheet1.xml").decode("utf-8")
            assert "<v />" not in sheet_text
            assert "<v></v>" not in sheet_text

    def test_excel_generator_preserves_formulas_and_borders(self, tmp_path: Path, sample_photos: dict):
        """Memverifikasi bahwa file template yang memiliki rumus & border sel tetap utuh setelah generate."""
        from openpyxl.styles import Border, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "LaporanBTS"

        thin = Side(border_style="thin", color="FF0000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        # Buat tabel dengan border dan rumus
        ws["A1"] = "Item"
        ws["B1"] = "Jumlah"
        ws["A1"].border = border
        ws["B1"].border = border

        ws["A2"] = "Perangkat 1"
        ws["B2"] = 10
        ws["A2"].border = border
        ws["B2"].border = border

        ws["A3"] = "Perangkat 2"
        ws["B3"] = 20
        ws["A3"].border = border
        ws["B3"].border = border

        ws["A4"] = "Total"
        ws["B4"] = "=SUM(B2:B3)"
        ws["A4"].border = border
        ws["B4"].border = border

        # Cell anchor untuk foto di B6
        ws.column_dimensions["B"].width = 25.0
        ws.row_dimensions[6].height = 40.0

        tpl_file = str(tmp_path / "template_with_borders.xlsx")
        wb.save(tpl_file)
        wb.close()

        out_file = str(tmp_path / "out_preserved.xlsx")
        field = PhotoFieldConfig(
            field_id="p1",
            label="Foto Site",
            sheet_name="LaporanBTS",
            cell_anchor="B6",
        )

        thread = ExcelGeneratorThread(
            template=tpl_file,
            output_path=out_file,
            fields=[field],
            photo_assignments={"p1": sample_photos["p1"]},
        )
        thread.run()

        assert os.path.isfile(out_file)

        # Periksa workbook hasil
        wb_res = openpyxl.load_workbook(out_file)
        ws_res = wb_res["LaporanBTS"]

        # 1. Verifikasi border tetap ada
        assert ws_res["A1"].border.left.style == "thin"
        assert ws_res["B1"].border.right.style == "thin"
        assert ws_res["A4"].border.top.style == "thin"
        assert ws_res["B4"].border.bottom.style == "thin"

        # 2. Verifikasi rumus tetap ada
        assert ws_res["B4"].value == "=SUM(B2:B3)"

        # 3. Verifikasi calculation properties terpasang
        if hasattr(wb_res, "calculation"):
            assert wb_res.calculation.fullCalcOnLoad is True
            assert wb_res.calculation.calcMode == "auto"

        wb_res.close()


class TestFormatOutputFilename:
    def test_default_fallback_when_empty_pattern(self):
        proj = ProjectData(project_id="p1", site_name="Site Malang 01", template_id="t1")
        tpl = TemplateConfig(template_id="t1", name="Template CME", template_file_path="")

        fn = format_output_filename("", project_data=proj, template=tpl)
        assert fn.startswith("Laporan_Site_Malang_01_")
        assert fn.endswith(".xlsx")

    def test_standard_variables_substitution(self):
        proj = ProjectData(project_id="p1", site_name="JKT-42", template_id="t1")
        tpl = TemplateConfig(template_id="t1", name="ATP Telkomsel", template_file_path="")

        fn = format_output_filename(
            "Laporan_{site_name}_{template_name}_{date}",
            project_data=proj,
            template=tpl,
        )
        assert fn.startswith("Laporan_JKT-42_ATP_Telkomsel_")
        assert fn.endswith(".xlsx")

    def test_field_variables_substitution(self):
        field1 = PhotoFieldConfig(
            field_id="fld_nama_site",
            label="Nama Site",
            sheet_name="S1",
            cell_anchor="A1",
            field_type="text",
        )
        field2 = PhotoFieldConfig(
            field_id="fld_spk",
            label="Nomor SPK",
            sheet_name="S1",
            cell_anchor="A2",
            field_type="text",
        )
        tpl = TemplateConfig(
            template_id="t1",
            name="Template CME",
            template_file_path="",
            fields=[field1, field2],
        )
        proj = ProjectData(
            project_id="p1",
            site_name="SITE-001",
            template_id="t1",
            text_assignments={"fld_nama_site": "Pasar Minggu", "fld_spk": "SPK-9988"},
        )

        fn = format_output_filename(
            "{template_name}_{Nama Site}_{Nomor SPK}_{date}",
            project_data=proj,
            template=tpl,
        )
        assert "Template_CME_Pasar_Minggu_SPK-9988_" in fn
        assert fn.endswith(".xlsx")

    def test_illegal_characters_and_whitespace_cleaning(self):
        proj = ProjectData(project_id="p1", site_name='JKT/01:Telkom*Sel?"<>|', template_id="t1")
        tpl = TemplateConfig(template_id="t1", name="Doc/CME", template_file_path="")

        fn = format_output_filename(
            "Doc_{site_name}",
            project_data=proj,
            template=tpl,
        )
        assert "/" not in fn
        assert ":" not in fn
        assert "*" not in fn
        assert "?" not in fn
        assert '"' not in fn
        assert "<" not in fn
        assert ">" not in fn
        assert "|" not in fn
        assert fn.endswith(".xlsx")


