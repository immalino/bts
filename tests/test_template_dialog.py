"""
Photo Rename Studio - Pengujian Unit Template Dialog & Excel Inspector (Sesi 10)

Memverifikasi:
1. Validasi form FieldEditDialog (label kosong, sheet kosong, koordinat sel tidak valid, ukuran custom).
2. Auto-load nama sheet dari file .xlsx ke TemplateEditorDialog.
3. Operasi visual tabel field foto (tambah, edit, hapus, naik, turun).
4. Validasi form TemplateEditorDialog (nama kosong, file tidak ada/tidak valid, tanpa field, koordinat rusak).
5. Integrasi persistensi template dengan StorageManager.
6. TemplateManagerDialog (daftar template, pemilihan template aktif, hapus template).
7. Fungsi utilitas excel_utils (validasi koordinat sel & pembaca nama sheet).
"""

from pathlib import Path
import openpyxl
import pytest
from PySide6.QtWidgets import QApplication, QDialog, QMessageBox

from src.core.models import PhotoFieldConfig, TemplateConfig
from src.core.storage import StorageManager
from src.ui.document_tabs import AddDocumentDialog
from src.ui.project_dialog import NewProjectDialog, ProjectManagerDialog
from src.ui.template_dialog import (
    FieldEditDialog,
    TemplateEditorDialog,
    TemplateManagerDialog,
)
from src.utils.excel_utils import get_sheet_names_from_excel, is_valid_cell_coordinate


@pytest.fixture(scope="session")
def qapp():
    """Fixture instance QApplication."""
    app = QApplication.instance()
    if app is None:
        app = QApplication([])
    return app


@pytest.fixture
def sample_xlsx(tmp_path: Path) -> str:
    """Fixture membuat file template Excel dummy dengan multi-sheet."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Antenna"
    ws2 = wb.create_sheet(title="Equipment")
    ws3 = wb.create_sheet(title="Tower")

    file_path = str(tmp_path / "test_template.xlsx")
    wb.save(file_path)
    wb.close()
    return file_path


# =====================================================================
# 1. Pengujian Excel Utils Helper (Koordinat Sel & Sheet Reader)
# =====================================================================

def test_is_valid_cell_coordinate():
    """Memverifikasi fungsi validasi koordinat sel Excel."""
    # Valid coordinates (single cell & cell range)
    assert is_valid_cell_coordinate("A1") is True
    assert is_valid_cell_coordinate("B5") is True
    assert is_valid_cell_coordinate("z100") is True
    assert is_valid_cell_coordinate("AA1") is True
    assert is_valid_cell_coordinate("$B$5") is True
    assert is_valid_cell_coordinate("  C10  ") is True
    assert is_valid_cell_coordinate("XFD1048576") is True
    assert is_valid_cell_coordinate("B1:D4") is True
    assert is_valid_cell_coordinate("a1:c10") is True
    assert is_valid_cell_coordinate("$B$1:$D$4") is True
    assert is_valid_cell_coordinate(" B1 : D4 ") is True

    # Invalid coordinates
    assert is_valid_cell_coordinate("") is False
    assert is_valid_cell_coordinate("   ") is False
    assert is_valid_cell_coordinate("123") is False
    assert is_valid_cell_coordinate("B") is False
    assert is_valid_cell_coordinate("5B") is False
    assert is_valid_cell_coordinate("A0") is False
    assert is_valid_cell_coordinate("ABCD1") is False  # Kolom lebih dari 3 huruf
    assert is_valid_cell_coordinate("A9999999") is False  # Melebihi batas baris Excel
    assert is_valid_cell_coordinate("B1:D4:E5") is False
    assert is_valid_cell_coordinate("B1:") is False
    assert is_valid_cell_coordinate(":D4") is False
    assert is_valid_cell_coordinate(None) is False  # type: ignore


def test_get_sheet_names_from_excel(sample_xlsx: str, tmp_path: Path):
    """Memverifikasi pembacaan daftar sheet dari berkas Excel."""
    sheets = get_sheet_names_from_excel(sample_xlsx)
    assert sheets == ["Antenna", "Equipment", "Tower"]

    # File tidak ada
    with pytest.raises(FileNotFoundError):
        get_sheet_names_from_excel(tmp_path / "non_existent.xlsx")

    # File bukan format Excel
    txt_file = tmp_path / "dummy.txt"
    txt_file.write_text("hello")
    with pytest.raises(ValueError):
        get_sheet_names_from_excel(txt_file)


# =====================================================================
# 2. Pengujian FieldEditDialog & Validasi Form
# =====================================================================

def test_field_edit_dialog_validation_empty_label(qapp, monkeypatch):
    """Memverifikasi validasi label kosong pada FieldEditDialog."""
    warning_called = False

    def mock_warning(*args, **kwargs):
        nonlocal warning_called
        warning_called = True

    monkeypatch.setattr(QMessageBox, "warning", mock_warning)

    dialog = FieldEditDialog(available_sheets=["Antenna", "Tower"])
    dialog.edit_label.setText("")
    dialog.combo_sheet.setCurrentText("Antenna")
    dialog.edit_cell.setText("B5")

    dialog._on_save_clicked()
    assert warning_called is True
    assert dialog.get_field_config() is None
    dialog.close()


def test_field_edit_dialog_validation_empty_sheet(qapp, monkeypatch):
    """Memverifikasi validasi sheet target kosong pada FieldEditDialog."""
    warning_called = False

    def mock_warning(*args, **kwargs):
        nonlocal warning_called
        warning_called = True

    monkeypatch.setattr(QMessageBox, "warning", mock_warning)

    dialog = FieldEditDialog()
    dialog.edit_label.setText("Foto Sektor 1")
    dialog.combo_sheet.setCurrentText("")
    dialog.edit_cell.setText("B5")

    dialog._on_save_clicked()
    assert warning_called is True
    assert dialog.get_field_config() is None
    dialog.close()


def test_field_edit_dialog_validation_invalid_cell(qapp, monkeypatch):
    """Memverifikasi validasi koordinat sel anchor tidak valid pada FieldEditDialog."""
    warning_called = False

    def mock_warning(*args, **kwargs):
        nonlocal warning_called
        warning_called = True

    monkeypatch.setattr(QMessageBox, "warning", mock_warning)

    dialog = FieldEditDialog(available_sheets=["Antenna"])
    dialog.edit_label.setText("Foto Sektor 1")
    dialog.combo_sheet.setCurrentText("Antenna")
    dialog.edit_cell.setText("SALAH_SEL")

    dialog._on_save_clicked()
    assert warning_called is True
    assert dialog.get_field_config() is None
    dialog.close()


def test_field_edit_dialog_save_valid(qapp, monkeypatch):
    """Memverifikasi penyimpanan sukses pada FieldEditDialog dengan mode fit_to_cell dan custom."""
    # 1. Mode fit_to_cell
    dialog = FieldEditDialog(available_sheets=["Antenna", "Tower"])
    dialog.edit_label.setText("Foto Antenna Sektor 1")
    dialog.combo_sheet.setCurrentText("Antenna")
    dialog.edit_cell.setText("B5")
    dialog.radio_fit.setChecked(True)
    dialog.spin_padding.setValue(3)

    dialog._on_save_clicked()
    res = dialog.get_field_config()
    assert res is not None
    assert res.label == "Foto Antenna Sektor 1"
    assert res.sheet_name == "Antenna"
    assert res.cell_anchor == "B5"
    assert res.sizing_mode == "fit_to_cell"
    assert res.padding == 3
    assert res.alignment == "center"
    dialog.close()

    # 2. Mode custom sizing dengan top-left alignment
    dialog2 = FieldEditDialog()
    dialog2.edit_label.setText("Foto Perangkat")
    dialog2.combo_sheet.setCurrentText("Equipment")
    dialog2.edit_cell.setText("C10")
    dialog2.radio_custom.setChecked(True)
    dialog2.radio_align_topleft.setChecked(True)
    dialog2.spin_width.setValue(450)
    dialog2.spin_height.setValue(350)
    dialog2.spin_padding.setValue(4)

    dialog2._on_save_clicked()
    res2 = dialog2.get_field_config()
    assert res2 is not None
    assert res2.label == "Foto Perangkat"
    assert res2.sizing_mode == "custom"
    assert res2.custom_width == 450
    assert res2.custom_height == 350
    assert res2.padding == 4
    assert res2.alignment == "top-left"
    dialog2.close()


def test_field_edit_dialog_load_existing_config(qapp):
    """Memverifikasi form FieldEditDialog terisi otomatis saat mengedit field yang sudah ada."""
    existing = PhotoFieldConfig(
        field_id="fld_existing_1",
        label="Foto Tower Monopole",
        sheet_name="Tower",
        cell_anchor="D12",
        sizing_mode="custom",
        custom_width=320,
        custom_height=240,
        padding=5,
    )

    dialog = FieldEditDialog(
        field_config=existing,
        available_sheets=["Antenna", "Tower", "Equipment"],
    )
    assert dialog.edit_label.text() == "Foto Tower Monopole"
    assert dialog.combo_sheet.currentText() == "Tower"
    assert dialog.edit_cell.text() == "D12"
    assert dialog.radio_custom.isChecked() is True
    assert dialog.spin_width.value() == 320
    assert dialog.spin_height.value() == 240
    assert dialog.spin_padding.value() == 5

    dialog._on_save_clicked()
    saved = dialog.get_field_config()
    assert saved is not None
    assert saved.field_id == "fld_existing_1"  # ID tetap terjaga
    dialog.close()


# =====================================================================
# 3. Pengujian TemplateEditorDialog (Auto-load Sheet & Visual Field Table)
# =====================================================================

def test_template_editor_auto_load_sheets(qapp, sample_xlsx: str):
    """Memverifikasi auto-load nama sheet saat file .xlsx disetel."""
    editor = TemplateEditorDialog()
    assert len(editor.available_sheets) == 0

    editor.set_template_file(sample_xlsx)
    assert editor.available_sheets == ["Antenna", "Equipment", "Tower"]
    assert "Sheet terdeteksi (3)" in editor.lbl_sheet_status.text()

    # File invalid
    editor.set_template_file("/invalid/path/non_existent.xlsx")
    assert len(editor.available_sheets) == 0
    assert "Gagal membaca sheet" in editor.lbl_sheet_status.text()
    editor.close()


def test_template_editor_validation(qapp, sample_xlsx: str, tmp_path: Path, monkeypatch):
    """Memverifikasi validasi form TemplateEditorDialog untuk seluruh kriteria error."""
    monkeypatch.setattr(QMessageBox, "warning", lambda *args, **kwargs: None)

    editor = TemplateEditorDialog()

    # 1. Nama template kosong
    editor.edit_name.setText("")
    editor.edit_file_path.setText(sample_xlsx)
    assert editor.validate_form() is False

    # 2. File template kosong
    editor.edit_name.setText("Template BTS")
    editor.edit_file_path.setText("")
    assert editor.validate_form() is False

    # 3. File template tidak ditemukan
    editor.edit_file_path.setText(str(tmp_path / "not_found.xlsx"))
    assert editor.validate_form() is False

    # 4. File template bukan .xlsx
    fake_txt = tmp_path / "sample.txt"
    fake_txt.write_text("bukan excel")
    editor.edit_file_path.setText(str(fake_txt))
    assert editor.validate_form() is False

    # 5. Tidak ada field foto (0 field)
    editor.edit_file_path.setText(sample_xlsx)
    editor.fields.clear()
    assert editor.validate_form() is False

    # 6. Field foto dengan koordinat sel rusak
    broken_field = PhotoFieldConfig(
        field_id="fld_bad",
        label="Foto Bad",
        sheet_name="Antenna",
        cell_anchor="SALAH_KOORDINAT",
    )
    editor.fields = [broken_field]
    assert editor.validate_form() is False

    # 7. Form valid sempurna
    valid_field = PhotoFieldConfig(
        field_id="fld_good",
        label="Foto Good",
        sheet_name="Antenna",
        cell_anchor="B5",
    )
    editor.fields = [valid_field]
    assert editor.validate_form() is True
    editor.close()


def test_template_editor_fields_crud_and_order(qapp, sample_xlsx: str, monkeypatch):
    """Memverifikasi visual table operasi: Tambah, Edit, Hapus, Naik, Turun field."""
    editor = TemplateEditorDialog()
    editor.set_template_file(sample_xlsx)

    # Tambah Field 1 via dialog mock
    mock_field_1 = PhotoFieldConfig(
        field_id="f1",
        label="Field Satu",
        sheet_name="Antenna",
        cell_anchor="B5",
    )
    monkeypatch.setattr(FieldEditDialog, "exec", lambda self: QDialog.DialogCode.Accepted)
    monkeypatch.setattr(FieldEditDialog, "get_field_config", lambda self: mock_field_1)

    editor._on_add_field_clicked()
    assert len(editor.fields) == 1
    assert editor.table_fields.rowCount() == 1
    assert editor.table_fields.item(0, 1).text() == "Field Satu"

    # Tambah Field 2
    mock_field_2 = PhotoFieldConfig(
        field_id="f2",
        label="Field Dua",
        sheet_name="Equipment",
        cell_anchor="C10",
    )
    monkeypatch.setattr(FieldEditDialog, "get_field_config", lambda self: mock_field_2)
    editor._on_add_field_clicked()
    assert len(editor.fields) == 2
    assert editor.table_fields.item(1, 1).text() == "Field Dua"

    # Edit Field 1
    editor.table_fields.selectRow(0)
    mock_field_1_edited = PhotoFieldConfig(
        field_id="f1",
        label="Field Satu Diedit",
        sheet_name="Antenna",
        cell_anchor="B6",
    )
    monkeypatch.setattr(FieldEditDialog, "get_field_config", lambda self: mock_field_1_edited)
    editor._on_edit_field_clicked()
    assert editor.fields[0].label == "Field Satu Diedit"
    assert editor.table_fields.item(0, 1).text() == "Field Satu Diedit"

    # Pindah Turun (baris 0 turun ke baris 1)
    editor.table_fields.selectRow(0)
    editor._on_move_down_clicked()
    assert editor.fields[0].label == "Field Dua"
    assert editor.fields[1].label == "Field Satu Diedit"

    # Pindah Naik (baris 1 naik ke baris 0)
    editor.table_fields.selectRow(1)
    editor._on_move_up_clicked()
    assert editor.fields[0].label == "Field Satu Diedit"
    assert editor.fields[1].label == "Field Dua"

    # Hapus Field
    monkeypatch.setattr(QMessageBox, "question", lambda *args, **kwargs: QMessageBox.StandardButton.Yes)
    editor.table_fields.selectRow(1)
    editor._on_delete_field_clicked()
    assert len(editor.fields) == 1
    assert editor.fields[0].label == "Field Satu Diedit"
    assert editor.table_fields.rowCount() == 1

    editor.close()


def test_template_editor_save_to_storage(qapp, sample_xlsx: str, tmp_path: Path):
    """Memverifikasi integrasi penyimpanan konfigurasi template ke StorageManager."""
    storage = StorageManager(base_dir=tmp_path / "storage")

    editor = TemplateEditorDialog(storage=storage)
    editor.edit_name.setText("Template Uji Coba")
    editor.set_template_file(sample_xlsx)
    editor.fields = [
        PhotoFieldConfig(
            field_id="fld_1",
            label="Foto 1",
            sheet_name="Antenna",
            cell_anchor="B5",
        )
    ]

    editor._on_save_template_clicked()
    saved = editor.get_template_config()
    assert saved is not None
    assert saved.name == "Template Uji Coba"

    # Pastikan tersimpan permanen di StorageManager
    loaded_list = storage.list_templates()
    assert len(loaded_list) == 1
    assert loaded_list[0].template_id == saved.template_id
    assert loaded_list[0].name == "Template Uji Coba"
    assert len(loaded_list[0].fields) == 1
    editor.close()


# =====================================================================
# 4. Pengujian TemplateManagerDialog (Daftar, Pilih, Hapus)
# =====================================================================

def test_template_manager_dialog_workflow(qapp, sample_xlsx: str, tmp_path: Path, monkeypatch):
    """Memverifikasi alur kerja TemplateManagerDialog: memuat, memilih, dan menghapus template."""
    storage = StorageManager(base_dir=tmp_path / "storage_manager")

    # Buat 2 template dummy di storage
    tpl1 = TemplateConfig(
        template_id="tpl_001",
        name="Template Alfa",
        template_file_path=sample_xlsx,
        fields=[PhotoFieldConfig("f1", "Foto A", "Antenna", "B5")],
    )
    tpl2 = TemplateConfig(
        template_id="tpl_002",
        name="Template Beta",
        template_file_path=sample_xlsx,
        fields=[
            PhotoFieldConfig("f2", "Foto B1", "Equipment", "C5"),
            PhotoFieldConfig("f3", "Foto B2", "Tower", "D10"),
        ],
    )
    storage.save_template(tpl1)
    storage.save_template(tpl2)

    manager = TemplateManagerDialog(storage=storage)
    assert manager.table_templates.rowCount() == 2
    assert manager.table_templates.item(0, 0).text() == "Template Alfa"
    assert manager.table_templates.item(1, 0).text() == "Template Beta"
    assert manager.btn_select.isEnabled() is False

    # Pilih baris 1 (Template Beta)
    manager.table_templates.selectRow(1)
    assert manager.btn_select.isEnabled() is True
    assert manager.btn_edit.isEnabled() is True
    assert manager.btn_delete.isEnabled() is True

    # Klik "Pilih Template Aktif"
    manager._on_select_template_clicked()
    selected = manager.get_selected_template()
    assert selected is not None
    assert selected.template_id == "tpl_002"
    assert selected.name == "Template Beta"

    # Test Hapus Template
    monkeypatch.setattr(QMessageBox, "question", lambda *args, **kwargs: QMessageBox.StandardButton.Yes)
    manager.table_templates.selectRow(0)
    manager._on_delete_template_clicked()

    # Template Alfa terhapus dari storage dan tabel
    assert len(storage.list_templates()) == 1
    assert manager.table_templates.rowCount() == 1
    assert manager.table_templates.item(0, 0).text() == "Template Beta"

    manager.close()


def test_field_edit_dialog_text_type_flow(qapp):
    """Memverifikasi pembuatan dan pengeditan field tipe teks di FieldEditDialog."""
    dlg = FieldEditDialog(available_sheets=["Site_Info", "Antenna"])
    assert dlg.radio_type_photo.isChecked() is True
    assert dlg.text_options_container.isHidden() is True

    # Pilih tipe teks
    dlg.radio_type_text.setChecked(True)
    assert dlg.photo_options_container.isHidden() is True
    assert dlg.text_options_container.isHidden() is False
    assert dlg.lbl_field_label.text() == "Label Field Teks:"

    dlg.edit_label.setText("Nama Teknisi Lead")
    dlg.combo_sheet.setCurrentText("Site_Info")
    dlg.edit_cell.setText("C2")
    dlg.edit_default_text.setText("Budi Santoso")

    dlg._on_save_clicked()
    res = dlg.get_field_config()
    assert res is not None
    assert res.label == "Nama Teknisi Lead"
    assert res.field_type == "text"
    assert res.sheet_name == "Site_Info"
    assert res.cell_anchor == "C2"
    assert res.default_text == "Budi Santoso"
    dlg.close()


def test_template_editor_mixed_fields_table(qapp, sample_xlsx: str):
    """Memverifikasi tampilan tabel field di TemplateEditorDialog dengan tipe campuran Foto dan Teks."""
    editor = TemplateEditorDialog()
    editor.set_template_file(sample_xlsx)

    editor.fields = [
        PhotoFieldConfig("f1", "Foto Antenna 1", "Antenna", "B5", field_type="photo"),
        PhotoFieldConfig("t1", "Nama Site BTS", "Antenna", "A1", field_type="text", default_text="Site-JKT"),
    ]
    editor._refresh_fields_table()

    assert editor.table_fields.rowCount() == 2
    # Baris 0: Foto
    assert editor.table_fields.item(0, 1).text() == "Foto Antenna 1"
    assert editor.table_fields.item(0, 2).text() == "Foto"
    assert editor.table_fields.item(0, 5).text() == "Auto-Fit Sel"
    assert editor.table_fields.item(0, 6).text() == "2 px"

    # Baris 1: Teks
    assert editor.table_fields.item(1, 1).text() == "Nama Site BTS"
    assert editor.table_fields.item(1, 2).text() == "Teks"
    assert "Teks (Default: Site-JKT)" in editor.table_fields.item(1, 5).text()
    assert editor.table_fields.item(1, 6).text() == "-"
    assert editor.table_fields.item(1, 7).text() == "-"

    editor.close()


def test_template_editor_output_filename_pattern_save_and_load(qapp, sample_xlsx: str, tmp_path: Path):
    """Memverifikasi TemplateEditorDialog menyimpan dan memuat kembali pola nama file output."""
    storage = StorageManager(base_dir=tmp_path / "storage_test")

    # 1. Buat template baru dengan pola kustom
    editor = TemplateEditorDialog(storage=storage)
    editor.edit_name.setText("Template Pola Kustom")
    editor.edit_file_path.setText(sample_xlsx)
    editor.edit_output_filename_pattern.setText("Laporan_{template_name}_{site_name}_{date}")
    editor.fields = [
        PhotoFieldConfig("f1", "Foto Antenna", "Antenna", "B5", field_type="photo")
    ]

    editor._on_save_template_clicked()
    saved = editor.get_template_config()
    assert saved is not None
    assert saved.output_filename_pattern == "Laporan_{template_name}_{site_name}_{date}"
    editor.close()

    # 2. Muat kembali template yang tersimpan ke editor
    loaded = storage.load_template(saved.template_id)
    assert loaded.output_filename_pattern == "Laporan_{template_name}_{site_name}_{date}"

    editor2 = TemplateEditorDialog(template=loaded, storage=storage)
    assert editor2.edit_output_filename_pattern.text() == "Laporan_{template_name}_{site_name}_{date}"
    editor2.close()


def test_dialog_dimensions_minimum_size(qapp, tmp_path: Path):
    """Memverifikasi batas ukuran minimum seluruh window dialog tidak terlalu kecil (tidak gepeng)."""
    storage = StorageManager(base_dir=tmp_path / "storage_dim")

    dlg_new_proj = NewProjectDialog(storage=storage)
    assert dlg_new_proj.minimumWidth() >= 480
    assert dlg_new_proj.minimumHeight() >= 320
    dlg_new_proj.close()

    dlg_proj_mgr = ProjectManagerDialog(storage=storage)
    assert dlg_proj_mgr.minimumWidth() >= 700
    assert dlg_proj_mgr.minimumHeight() >= 450
    dlg_proj_mgr.close()

    dlg_add_doc = AddDocumentDialog(storage=storage)
    assert dlg_add_doc.minimumWidth() >= 450
    assert dlg_add_doc.minimumHeight() >= 280
    dlg_add_doc.close()

    dlg_tpl_mgr = TemplateManagerDialog(storage=storage)
    assert dlg_tpl_mgr.minimumWidth() >= 700
    assert dlg_tpl_mgr.minimumHeight() >= 450
    dlg_tpl_mgr.close()

    dlg_tpl_editor = TemplateEditorDialog(storage=storage)
    assert dlg_tpl_editor.minimumWidth() >= 700
    assert dlg_tpl_editor.minimumHeight() >= 500
    dlg_tpl_editor.close()

    dlg_field = FieldEditDialog()
    assert dlg_field.minimumWidth() >= 460
    assert dlg_field.minimumHeight() >= 500
    dlg_field.close()


def test_template_editor_move_top_and_bottom(qapp, sample_xlsx: str):
    """Memverifikasi pemindahan urutan field ke paling atas dan paling bawah."""
    editor = TemplateEditorDialog()
    editor.set_template_file(sample_xlsx)

    f1 = PhotoFieldConfig("f1", "Field 1", "Sheet1", "A1")
    f2 = PhotoFieldConfig("f2", "Field 2", "Sheet1", "A2")
    f3 = PhotoFieldConfig("f3", "Field 3", "Sheet1", "A3")
    f4 = PhotoFieldConfig("f4", "Field 4", "Sheet1", "A4")
    editor.fields = [f1, f2, f3, f4]
    editor._refresh_fields_table()

    # Pindahkan f3 (index 2) ke paling atas
    editor.table_fields.selectRow(2)
    editor._on_move_top_clicked()
    assert [f.field_id for f in editor.fields] == ["f3", "f1", "f2", "f4"]
    assert editor.table_fields.currentRow() == 0

    # Pindahkan f1 (index 1) ke paling bawah
    editor.table_fields.selectRow(1)
    editor._on_move_bottom_clicked()
    assert [f.field_id for f in editor.fields] == ["f3", "f2", "f4", "f1"]
    assert editor.table_fields.currentRow() == 3

    # Naik 1 tingkat (f1 dari index 3 ke 2)
    editor._on_move_up_clicked()
    assert [f.field_id for f in editor.fields] == ["f3", "f2", "f1", "f4"]
    assert editor.table_fields.currentRow() == 2

    # Turun 1 tingkat (f2 dari index 1 ke 2)
    editor.table_fields.selectRow(1)
    editor._on_move_down_clicked()
    assert [f.field_id for f in editor.fields] == ["f3", "f1", "f2", "f4"]
    assert editor.table_fields.currentRow() == 2

    editor.close()


def test_template_editor_duplicate_field(qapp, sample_xlsx: str):
    """Memverifikasi duplikasi field 1-klik."""
    editor = TemplateEditorDialog()
    editor.set_template_file(sample_xlsx)

    f1 = PhotoFieldConfig("f1", "Foto Antenna Sector 1", "Antenna", "B5", sizing_mode="fit_to_cell")
    editor.fields = [f1]
    editor._refresh_fields_table()

    editor.table_fields.selectRow(0)
    editor._on_duplicate_field_clicked()

    assert len(editor.fields) == 2
    dup = editor.fields[1]
    assert dup.field_id != f1.field_id
    assert dup.label == "Foto Antenna Sector 1 (Salinan)"
    assert dup.sheet_name == "Antenna"
    assert dup.cell_anchor == "B5"
    assert dup.sizing_mode == "fit_to_cell"
    assert editor.table_fields.currentRow() == 1
    editor.close()


def test_template_editor_inline_editing(qapp, sample_xlsx: str):
    """Memverifikasi in-place editing langsung pada tabel field tanpa membuka modal dialog."""
    editor = TemplateEditorDialog()
    editor.set_template_file(sample_xlsx)

    f1 = PhotoFieldConfig("f1", "Label Awal", "Antenna", "B5")
    editor.fields = [f1]
    editor._refresh_fields_table()

    # Ubah label via cell widget
    edit_label = editor.table_fields.cellWidget(0, 1)
    assert edit_label is not None
    edit_label.setText("Label Diedit Langsung")
    assert editor.fields[0].label == "Label Diedit Langsung"
    assert editor.table_fields.item(0, 1).text() == "Label Diedit Langsung"

    # Ubah sheet target via cell widget
    combo_sheet = editor.table_fields.cellWidget(0, 3)
    assert combo_sheet is not None
    combo_sheet.setCurrentText("Equipment")
    assert editor.fields[0].sheet_name == "Equipment"
    assert editor.table_fields.item(0, 3).text() == "Equipment"

    # Ubah sel anchor via cell widget
    edit_cell = editor.table_fields.cellWidget(0, 4)
    assert edit_cell is not None
    edit_cell.setText("D12")
    assert editor.fields[0].cell_anchor == "D12"
    assert editor.table_fields.item(0, 4).text() == "D12"

    editor.close()


def test_template_editor_export_and_import_json(qapp, sample_xlsx: str, tmp_path: Path, monkeypatch):
    """Memverifikasi fitur Export dan Import field dari/ke file JSON (mode Replace dan Append)."""
    import json
    from PySide6.QtWidgets import QFileDialog

    editor = TemplateEditorDialog()
    editor.set_template_file(sample_xlsx)

    f_photo = PhotoFieldConfig("f_p", "Foto Antenna", "Antenna", "B5", field_type="photo", sizing_mode="fit_to_cell")
    f_text = PhotoFieldConfig("f_t", "Nama Site", "Antenna", "A1", field_type="text", default_text="JKT-01")
    editor.fields = [f_photo, f_text]
    editor._refresh_fields_table()

    json_export_path = str(tmp_path / "fields_exported.json")

    # Mocks agar modal dialog QMessageBox tidak memblokir jalannya pengujian
    monkeypatch.setattr(QMessageBox, "information", lambda *args, **kwargs: QMessageBox.StandardButton.Ok)
    monkeypatch.setattr(QMessageBox, "warning", lambda *args, **kwargs: QMessageBox.StandardButton.Ok)

    # 1. Test Export JSON
    monkeypatch.setattr(
        QFileDialog,
        "getSaveFileName",
        lambda *args, **kwargs: (json_export_path, "JSON Files (*.json)"),
    )
    editor._on_export_json_clicked()
    assert Path(json_export_path).is_file()

    with open(json_export_path, "r", encoding="utf-8") as fp:
        exported_data = json.load(fp)
    fields_exported = exported_data.get("fields", exported_data)
    assert len(fields_exported) == 2
    assert fields_exported[0]["label"] == "Foto Antenna"
    assert fields_exported[1]["label"] == "Nama Site"

    # 2. Test Import JSON - Mode Replace ("Ganti Semua")
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileName",
        lambda *args, **kwargs: (json_export_path, "JSON Files (*.json)"),
    )

    def mock_exec_replace(self):
        for btn in self.buttons():
            if "Ganti Semua" in btn.text():
                self._clicked_btn = btn
                return 0
        return 0

    monkeypatch.setattr(QMessageBox, "exec", mock_exec_replace)
    monkeypatch.setattr(QMessageBox, "clickedButton", lambda self: getattr(self, "_clicked_btn", None))

    # Ubah fields editor ke dummy dulu
    editor.fields = [PhotoFieldConfig("dummy", "Dummy", "Sheet1", "A1")]
    editor._on_import_json_clicked()
    assert len(editor.fields) == 2
    assert editor.fields[0].label == "Foto Antenna"
    assert editor.fields[1].label == "Nama Site"

    # 3. Test Import JSON - Mode Append ("Tambahkan")
    def mock_exec_append(self):
        for btn in self.buttons():
            if "Tambahkan" in btn.text():
                self._clicked_btn = btn
                return 0
        return 0

    monkeypatch.setattr(QMessageBox, "exec", mock_exec_append)
    editor._on_import_json_clicked()
    assert len(editor.fields) == 4

    editor.close()




