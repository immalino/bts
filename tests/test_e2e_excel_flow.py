"""
Photo Rename Studio - Pengujian Integrasi End-to-End (E2E) Alur Generate Dokumen Excel (SESI 12)

Memverifikasi:
1. Alur lengkap end-to-end: Pilih Site -> Pilih Template -> Pasang Foto -> Klik Generate -> Dokumen Excel Terbentuk.
2. Keberadaan gambar tertanam di dalam workbook Excel hasil output.
3. Validasi kegagalan generate jika belum ada foto atau file template tidak ditemukan.
4. Pembatalan proses generate via cancel_export().
5. Dialog laporan hasil dengan tombol satu-klik 'Buka Dokumen Excel' (QDesktopServices.openUrl) dan 'Buka Folder Output'.
6. Pembaruan metadata project (last_exported_at) dan persistensi ke storage lokal.
"""

import os
from pathlib import Path
import openpyxl
from PIL import Image
import pytest
from PySide6.QtCore import QCoreApplication, QUrl
from PySide6.QtGui import QDesktopServices
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox

from src.core.models import PhotoFieldConfig, ProjectData, TemplateConfig
from src.core.storage import StorageManager
from src.ui.main_window import MainWindow
from src.ui.report_dialog import ReportDialog


@pytest.fixture
def dummy_excel_template(tmp_path: Path) -> str:
    """Fixture membuat file template Excel nyata dengan beberapa sheet."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Antenna"
    ws1["A1"] = "DOKUMENTASI FOTO SEKTORAL"
    ws1.column_dimensions["B"].width = 25.0
    ws1.row_dimensions[5].height = 40.0

    ws2 = wb.create_sheet(title="Equipment")
    ws2["A1"] = "DOKUMENTASI PERANGKAT BTS"
    ws2.column_dimensions["C"].width = 30.0
    ws2.row_dimensions[10].height = 45.0

    tpl_path = str(tmp_path / "template_bts_telkom.xlsx")
    wb.save(tpl_path)
    wb.close()
    return tpl_path


@pytest.fixture
def dummy_photos(tmp_path: Path) -> tuple[str, str]:
    """Fixture membuat 2 file gambar JPEG nyata."""
    p1 = str(tmp_path / "photo_sector1.jpg")
    p2 = str(tmp_path / "photo_rack.jpg")

    Image.new("RGB", (300, 200), color="blue").save(p1)
    Image.new("RGB", (400, 300), color="green").save(p2)
    return p1, p2


def test_e2e_excel_generation_full_success(
    qapp,
    dummy_excel_template: str,
    dummy_photos: tuple[str, str],
    tmp_path: Path,
    monkeypatch,
):
    """
    Skenario E2E 1: Alur Lengkap Sukses Pembuatan Dokumen Excel Laporan.
    - Site aktif disetel dengan template multi-sheet.
    - 2 foto dipasang pada baris field dinamis.
    - User menekan tombol 'Generate Dokumen Excel'.
    - Output Excel tersimpan di direktori tujuan dengan gambar tertanam.
    - ReportDialog muncul dengan tombol 'Buka Dokumen Excel'.
    """
    storage = StorageManager(base_dir=tmp_path / "storage_e2e")
    photo1, photo2 = dummy_photos

    fields = [
        PhotoFieldConfig(
            field_id="fld_ant",
            label="Foto Antena Sector 1",
            sheet_name="Antenna",
            cell_anchor="B5",
            sizing_mode="fit_to_cell",
        ),
        PhotoFieldConfig(
            field_id="fld_eq",
            label="Foto Rack Perangkat",
            sheet_name="Equipment",
            cell_anchor="C10",
            sizing_mode="custom",
            custom_width=350,
            custom_height=250,
        ),
    ]

    template = TemplateConfig(
        template_id="tpl_e2e",
        name="Template BTS E2E",
        template_file_path=dummy_excel_template,
        fields=fields,
    )
    storage.save_template(template)

    project = ProjectData(
        project_id="proj_e2e_01",
        site_name="JKT-999 Kebayoran",
        template_id="tpl_e2e",
        photo_assignments={},
    )
    storage.save_project(project)

    window = MainWindow(storage=storage)
    window.set_project(project)

    # Pasang 2 foto pada baris dinamis
    window.dynamic_rows[0].set_photo(photo1)
    window.dynamic_rows[1].set_photo(photo2)

    # Mock file save dialog untuk menentukan lokasi output dokumen
    out_file = str(tmp_path / "output_reports" / "Laporan_JKT_999_Hasil.xlsx")
    monkeypatch.setattr(
        QFileDialog,
        "getSaveFileName",
        lambda *args, **kwargs: (out_file, "Excel Spreadsheet (*.xlsx)"),
    )

    # Mock ReportDialog.exec agar dialog terintersepsi dan tidak memblok test runner
    opened_dialogs = []
    monkeypatch.setattr(ReportDialog, "exec", lambda self: opened_dialogs.append(self))

    # Trigger aksi ekspor / generate
    window.start_export()

    # Tunggu background worker thread selesai
    assert window.excel_generator_thread is not None
    window.excel_generator_thread.wait(10000)
    QCoreApplication.processEvents()

    # 1. Verifikasi file Excel output terbuat di disk
    assert os.path.isfile(out_file), f"File Excel tidak ditemukan pada {out_file}"

    # 2. Verifikasi konten dokumen Excel yang dihasilkan
    wb_result = openpyxl.load_workbook(out_file)
    assert "Antenna" in wb_result.sheetnames
    assert "Equipment" in wb_result.sheetnames
    ws_ant = wb_result["Antenna"]
    ws_eq = wb_result["Equipment"]

    # Verifikasi gambar tertanam di kedua sheet
    assert len(ws_ant._images) >= 1
    assert len(ws_eq._images) >= 1
    wb_result.close()

    # 3. Verifikasi ReportDialog
    assert len(opened_dialogs) == 1
    report = opened_dialogs[0]
    assert report.success_count == 2
    assert report.fail_count == 0
    assert report.excel_path == out_file

    # Verifikasi tombol 'Buka Dokumen Excel' tersedia dan dapat diklik
    assert hasattr(report, "btn_open_excel")
    assert report.btn_open_excel.isVisible() or report.btn_open_excel.isEnabled()

    opened_urls = []
    monkeypatch.setattr(QDesktopServices, "openUrl", lambda url: opened_urls.append(url))
    report._on_open_excel_clicked()
    assert len(opened_urls) == 1
    assert os.path.normpath(out_file) == os.path.normpath(opened_urls[0].toLocalFile())

    # 4. Verifikasi update last_exported_at pada data project
    reloaded_proj = storage.load_project(project.project_id)
    assert reloaded_proj.last_exported_at is not None

    window.close()


def test_e2e_excel_validation_no_photos(
    qapp,
    dummy_excel_template: str,
    tmp_path: Path,
    monkeypatch,
):
    """
    Skenario E2E 2: Validasi Gagal Jika Belum Ada Foto yang Dimasukkan.
    """
    storage = StorageManager(base_dir=tmp_path / "storage_val_1")
    template = TemplateConfig(
        template_id="tpl_val1",
        name="Template Uji 1",
        template_file_path=dummy_excel_template,
        fields=[PhotoFieldConfig("f1", "Foto 1", "Antenna", "B5")],
    )
    storage.save_template(template)
    project = ProjectData("p_val1", "Site Kosong", "tpl_val1")
    storage.save_project(project)

    window = MainWindow(storage=storage)
    window.set_project(project)

    warning_dialogs = []
    monkeypatch.setattr(QMessageBox, "warning", lambda parent, title, msg: warning_dialogs.append((title, msg)))

    # Eksekusi ekspor tanpa memasukkan foto
    window.start_export()

    assert len(warning_dialogs) == 1
    assert "Foto Belum Dimasukkan" in warning_dialogs[0][0]
    assert window.excel_generator_thread is None
    window.close()


def test_e2e_excel_validation_missing_template_file(
    qapp,
    dummy_photos: tuple[str, str],
    tmp_path: Path,
    monkeypatch,
):
    """
    Skenario E2E 3: Validasi Gagal Jika File Template Fisik Tidak Ditemukan di Disk.
    """
    storage = StorageManager(base_dir=tmp_path / "storage_val_2")
    photo1, _ = dummy_photos

    template = TemplateConfig(
        template_id="tpl_val2",
        name="Template Hilang",
        template_file_path=str(tmp_path / "file_hilang.xlsx"),
        fields=[PhotoFieldConfig("f1", "Foto 1", "Antenna", "B5")],
    )
    storage.save_template(template)
    project = ProjectData("p_val2", "Site X", "tpl_val2")
    storage.save_project(project)

    window = MainWindow(storage=storage)
    window.set_project(project)
    window.dynamic_rows[0].set_photo(photo1)

    warning_dialogs = []
    monkeypatch.setattr(QMessageBox, "warning", lambda parent, title, msg: warning_dialogs.append((title, msg)))

    window.start_export()

    assert len(warning_dialogs) == 1
    assert "Template Tidak Ditemukan" in warning_dialogs[0][0]
    assert window.excel_generator_thread is None
    window.close()


def test_e2e_excel_cancellation(
    qapp,
    dummy_excel_template: str,
    dummy_photos: tuple[str, str],
    tmp_path: Path,
    monkeypatch,
):
    """
    Skenario E2E 4: Pembatalan Ekspor Excel di Tengah Jalan.
    """
    storage = StorageManager(base_dir=tmp_path / "storage_cancel")
    photo1, photo2 = dummy_photos

    fields = [
        PhotoFieldConfig("f1", "Foto 1", "Antenna", "B5"),
        PhotoFieldConfig("f2", "Foto 2", "Equipment", "C10"),
    ]
    template = TemplateConfig(
        template_id="tpl_cancel",
        name="Template Cancel",
        template_file_path=dummy_excel_template,
        fields=fields,
    )
    storage.save_template(template)
    project = ProjectData("p_cancel", "Site Cancel", "tpl_cancel")
    storage.save_project(project)

    window = MainWindow(storage=storage)
    window.set_project(project)
    window.dynamic_rows[0].set_photo(photo1)
    window.dynamic_rows[1].set_photo(photo2)

    out_file = str(tmp_path / "output_cancelled.xlsx")
    monkeypatch.setattr(
        QFileDialog,
        "getSaveFileName",
        lambda *args, **kwargs: (out_file, "Excel Spreadsheet (*.xlsx)"),
    )

    info_dialogs = []
    monkeypatch.setattr(QMessageBox, "information", lambda parent, title, msg: info_dialogs.append((title, msg)))

    # Mulai generate lalu langsung batalkan
    window.start_export()
    assert window.excel_generator_thread is not None
    window.cancel_export()

    window.excel_generator_thread.wait(5000)
    QCoreApplication.processEvents()

    assert len(info_dialogs) == 1
    assert "Dibatalkan" in info_dialogs[0][0]
    # File output tidak boleh tersimpan karena proses dibatalkan
    assert not os.path.exists(out_file)
    window.close()


def test_e2e_mixed_photo_and_text_excel_generation(
    qapp,
    dummy_excel_template: str,
    dummy_photos: tuple[str, str],
    tmp_path: Path,
    monkeypatch,
):
    """
    Skenario E2E 5: Alur Lengkap Dokumen Template Campuran Foto dan Teks ke Excel.
    """
    storage = StorageManager(base_dir=tmp_path / "storage_mixed_e2e")
    photo1, _ = dummy_photos

    fields = [
        PhotoFieldConfig("f_ant", "Foto Antenna Sector 1", "Antenna", "B5", field_type="photo"),
        PhotoFieldConfig("t_site", "Nama Site BTS", "Antenna", "A1", field_type="text"),
        PhotoFieldConfig("t_eng", "Nama Engineer", "Equipment", "B2", field_type="text", default_text="Budi (Lead)"),
    ]
    template = TemplateConfig(
        template_id="tpl_mixed_e2e",
        name="Template Telkom Foto & Teks",
        template_file_path=dummy_excel_template,
        fields=fields,
    )
    storage.save_template(template)

    project = ProjectData("p_mixed_e2e", "JKT-042 Pasar Minggu", "tpl_mixed_e2e")
    storage.save_project(project)

    window = MainWindow(storage=storage)
    window.set_project(project)

    # 1. Pasang foto ke field foto
    window.dynamic_rows[0].set_photo(photo1)
    # 2. Ketik teks ke field teks
    window.dynamic_rows[1].set_text("PASAR-MINGGU-042")
    # 3. Field teks ke-3 dibiarkan default ("Budi (Lead)")

    out_file = str(tmp_path / "output_mixed_final.xlsx")
    monkeypatch.setattr(
        QFileDialog,
        "getSaveFileName",
        lambda *args, **kwargs: (out_file, "Excel Spreadsheet (*.xlsx)"),
    )

    # Mock ReportDialog exec agar test tidak terblokir
    monkeypatch.setattr(ReportDialog, "exec", lambda self: QDialog.DialogCode.Accepted)

    # Eksekusi ekspor
    window.start_export()
    assert window.excel_generator_thread is not None
    window.excel_generator_thread.wait(5000)
    QCoreApplication.processEvents()

    assert os.path.exists(out_file)

    # Buka workbook dan verifikasi bahwa teks dan gambar berhasil disematkan
    wb = openpyxl.load_workbook(out_file)
    assert wb["Antenna"]["A1"].value == "PASAR-MINGGU-042"
    assert wb["Equipment"]["B2"].value == "Budi (Lead)"
    assert len(wb["Antenna"]._images) == 1
    wb.close()

    window.close()


