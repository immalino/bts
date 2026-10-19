"""
Photo Rename Studio - Excel Utilities
Fungsi pembantu untuk kalkulasi dimensi sel & merged cell di openpyxl
serta kalkulasi resize proporsional gambar (auto-fit cell).
"""

from datetime import datetime
import os
from pathlib import Path
import re
from typing import Any, Dict, List, Optional, Tuple, Union
import openpyxl
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import coordinate_to_tuple, get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.worksheet import Worksheet


def get_cell_dimensions(worksheet: Worksheet, cell_coordinate: str) -> Tuple[int, int]:
    """
    Menghitung dimensi lebar dan tinggi sel target dalam satuan piksel.
    Jika sel berada dalam rentang sel gabungan (merged cells), fungsi ini akan
    menjumlahkan seluruh lebar kolom dan tinggi baris dalam rentang tersebut.

    Konversi standar:
    - Lebar kolom (karakter openpyxl) -> piksel: round(lebar_char * 7 + 5)
    - Tinggi baris (points openpyxl) -> piksel: round(tinggi_pt * (96 / 72)) ≈ round(tinggi_pt * 1.333)

    Args:
        worksheet: Objek Worksheet dari openpyxl.
        cell_coordinate: Alamat koordinat sel (contoh: "B5", "C10").

    Returns:
        Tuple[int, int]: (lebar_px, tinggi_px) dalam piksel.
    """
    clean_coord = cell_coordinate.strip().replace("$", "").upper()
    if not clean_coord:
        raise ValueError(f"Koordinat sel tidak boleh kosong: '{cell_coordinate}'")

    # Cari apakah sel berada dalam rentang merged cells
    target_range = None
    for rng in worksheet.merged_cells.ranges:
        if clean_coord in rng:
            target_range = rng
            break

    if target_range is not None:
        min_col, min_row, max_col, max_row = target_range.bounds
    else:
        try:
            row, col = coordinate_to_tuple(clean_coord)
        except Exception as err:
            raise ValueError(f"Format koordinat sel tidak valid '{cell_coordinate}': {err}") from err
        min_col, max_col = col, col
        min_row, max_row = row, row

    # Hitung total lebar kolom dalam piksel
    total_width_px = 0
    for col_idx in range(min_col, max_col + 1):
        col_letter = get_column_letter(col_idx)
        col_dim = worksheet.column_dimensions.get(col_letter)
        if col_dim is not None and col_dim.width is not None:
            char_width = float(col_dim.width)
        elif getattr(worksheet.sheet_format, "defaultColWidth", None) is not None:
            char_width = float(worksheet.sheet_format.defaultColWidth)
        else:
            # Standar default kolom di Excel
            char_width = 8.43

        col_px = max(1, int(round(char_width * 7.0 + 5.0)))
        total_width_px += col_px

    # Hitung total tinggi baris dalam piksel
    total_height_px = 0
    for row_idx in range(min_row, max_row + 1):
        row_dim = worksheet.row_dimensions.get(row_idx)
        if row_dim is not None and row_dim.height is not None:
            pt_height = float(row_dim.height)
        elif getattr(worksheet.sheet_format, "defaultRowHeight", None) is not None:
            pt_height = float(worksheet.sheet_format.defaultRowHeight)
        else:
            # Standar default baris di Excel (15 pt = 20 px)
            pt_height = 15.0

        row_px = max(1, int(round(pt_height * (96.0 / 72.0))))
        total_height_px += row_px

    return (total_width_px, total_height_px)


def calculate_fitted_dimensions(
    img_width: int,
    img_height: int,
    max_width: int,
    max_height: int,
    padding: int = 2,
) -> Tuple[int, int]:
    """
    Menghitung ukuran dimensi gambar (lebar, tinggi) dalam piksel yang pas (fit)
    ke dalam batas kotak sel target dengan tetap mempertahankan rasio aspek asli (aspect ratio)
    serta menyisakan margin/padding di setiap sisinya agar tidak menabrak border sel.

    Args:
        img_width: Lebar asli gambar dalam piksel.
        img_height: Tinggi asli gambar dalam piksel.
        max_width: Lebar kotak sel target dalam piksel.
        max_height: Tinggi kotak sel target dalam piksel.
        padding: Margin internal di setiap sisi dalam piksel (default 2px).

    Returns:
        Tuple[int, int]: (lebar_fit, tinggi_fit) dalam piksel.
    """
    pad = max(0, padding)
    avail_w = max_width - (2 * pad)
    avail_h = max_height - (2 * pad)

    # Pastikan ruang tersedia minimal 1 piksel
    if avail_w <= 0:
        avail_w = max(1, max_width)
    if avail_h <= 0:
        avail_h = max(1, max_height)

    if img_width <= 0 or img_height <= 0:
        return (avail_w, avail_h)

    # Skala proporsional
    scale = min(avail_w / img_width, avail_h / img_height)
    new_w = max(1, int(round(img_width * scale)))
    new_h = max(1, int(round(img_height * scale)))

    return (new_w, new_h)


def calculate_image_anchor(
    worksheet: Worksheet,
    cell_coordinate: str,
    img_width: int,
    img_height: int,
    alignment: str = "center",
) -> OneCellAnchor:
    """
    Menghitung dan menghasilkan objek OneCellAnchor openpyxl untuk penempatan gambar pada worksheet.
    Mendukung perataan:
    - 'center': Gambar diposisikan persis di tengah (horizontal & vertikal) sel / merged cells.
    - 'top-left': Gambar diposisikan di sudut kiri-atas sel / merged cells (colOff=0, rowOff=0).

    Args:
        worksheet: Objek Worksheet dari openpyxl.
        cell_coordinate: Koordinat sel target (misal: "B5" atau rentang "B5:E15").
        img_width: Lebar gambar dalam piksel.
        img_height: Tinggi gambar dalam piksel.
        alignment: Mode perataan ('center' atau 'top-left'). Default 'center'.

    Returns:
        OneCellAnchor: Objek jangkar gambar dengan marker _from dan ukuran ext.
    """
    clean_coord = cell_coordinate.strip().replace("$", "").upper()
    if not clean_coord:
        raise ValueError(f"Koordinat sel tidak boleh kosong: '{cell_coordinate}'")

    # Cari apakah sel berada dalam rentang merged cells
    target_range = None
    for rng in worksheet.merged_cells.ranges:
        if clean_coord in rng:
            target_range = rng
            break

    if target_range is not None:
        min_col, min_row, max_col, max_row = target_range.bounds
    else:
        if ":" in clean_coord:
            c1, c2 = clean_coord.split(":", 1)
            r1, c1_idx = coordinate_to_tuple(c1)
            r2, c2_idx = coordinate_to_tuple(c2)
            min_col, max_col = min(c1_idx, c2_idx), max(c1_idx, c2_idx)
            min_row, max_row = min(r1, r2), max(r1, r2)
        else:
            try:
                row, col = coordinate_to_tuple(clean_coord)
            except Exception as err:
                raise ValueError(f"Format koordinat sel tidak valid '{cell_coordinate}': {err}") from err
            min_col, max_col = col, col
            min_row, max_row = row, row

    # Hitung lebar tiap kolom dalam piksel dan total lebar
    col_widths = []
    total_w = 0
    for col_idx in range(min_col, max_col + 1):
        col_letter = get_column_letter(col_idx)
        col_dim = worksheet.column_dimensions.get(col_letter)
        if col_dim is not None and col_dim.width is not None:
            char_w = float(col_dim.width)
        elif getattr(worksheet.sheet_format, "defaultColWidth", None) is not None:
            char_w = float(worksheet.sheet_format.defaultColWidth)
        else:
            char_w = 8.43
        col_px = max(1, int(round(char_w * 7.0 + 5.0)))
        col_widths.append((col_idx, col_px))
        total_w += col_px

    # Hitung tinggi tiap baris dalam piksel dan total tinggi
    row_heights = []
    total_h = 0
    for row_idx in range(min_row, max_row + 1):
        row_dim = worksheet.row_dimensions.get(row_idx)
        if row_dim is not None and row_dim.height is not None:
            pt_h = float(row_dim.height)
        elif getattr(worksheet.sheet_format, "defaultRowHeight", None) is not None:
            pt_h = float(worksheet.sheet_format.defaultRowHeight)
        else:
            pt_h = 15.0
        row_px = max(1, int(round(pt_h * (96.0 / 72.0))))
        row_heights.append((row_idx, row_px))
        total_h += row_px

    # Tentukan offset piksel
    if alignment.lower() == "top-left":
        offset_x_px = 0
        offset_y_px = 0
    else:
        # Default: center
        offset_x_px = max(0, (total_w - img_width) // 2)
        offset_y_px = max(0, (total_h - img_height) // 2)

    # Tentukan kolom mulai dan sisa colOff
    start_col = min_col
    curr_off_x = offset_x_px
    for c_idx, c_px in col_widths:
        if curr_off_x < c_px or c_idx == max_col:
            start_col = c_idx
            break
        curr_off_x -= c_px

    # Tentukan baris mulai dan sisa rowOff
    start_row = min_row
    curr_off_y = offset_y_px
    for r_idx, r_px in row_heights:
        if curr_off_y < r_px or r_idx == max_row:
            start_row = r_idx
            break
        curr_off_y -= r_px

    # Openpyxl AnchorMarker menggunakan indeks 0-based: col - 1, row - 1
    marker = AnchorMarker(
        col=start_col - 1,
        colOff=pixels_to_EMU(curr_off_x),
        row=start_row - 1,
        rowOff=pixels_to_EMU(curr_off_y),
    )
    size = XDRPositiveSize2D(pixels_to_EMU(img_width), pixels_to_EMU(img_height))
    return OneCellAnchor(_from=marker, ext=size)


def is_valid_cell_coordinate(cell_coordinate: str) -> bool:
    """
    Validasi format koordinat sel Excel (misal: 'A1', 'B5', 'AA10').
    Mengembalikan True jika format koordinat valid, False jika tidak.
    """
    if not isinstance(cell_coordinate, str):
        return False
    clean = cell_coordinate.strip().replace("$", "").upper()
    if not clean:
        return False
    import re
    if not re.match(r"^[A-Z]{1,3}[1-9][0-9]{0,6}$", clean):
        return False
    try:
        row, col = coordinate_to_tuple(clean)
        return 1 <= row <= 1048576 and 1 <= col <= 16384
    except Exception:
        return False


def get_sheet_names_from_excel(file_path: Union[str, Path]) -> List[str]:
    """
    Membaca daftar nama sheet dari file workbook Excel (.xlsx, .xlsm, dll) secara cepat.

    Args:
        file_path: Path ke file Excel.

    Returns:
        List[str]: Daftar nama sheet di workbook.

    Raises:
        FileNotFoundError: Jika file tidak ditemukan.
        ValueError: Jika format file bukan Excel atau korup.
    """
    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"File Excel tidak ditemukan: {file_path}")
    if path.suffix.lower() not in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        raise ValueError(f"File bukan format Excel spreadsheet yang didukung (.xlsx): {file_path}")

    try:
        wb = openpyxl.load_workbook(str(path), read_only=True)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()
    except Exception as err:
        if isinstance(err, (FileNotFoundError, ValueError)):
            raise
        raise ValueError(f"Gagal membaca sheet dari file Excel '{file_path}': {err}") from err


def sanitize_excel_output(file_path: str) -> None:
    """
    Membersihkan dan menstandarkan arsip berkas Excel (.xlsx) pasca penyimpanan openpyxl:
    1. Menghapus 'xl/calcChain.xml' usang dari arsip dan referensinya di workbook.xml.rels
       serta [Content_Types].xml. Hal ini mencegah peringatan korupsi kalkulasi rantai Excel
       ('We found a problem with some content... Removed Records: Formula from /xl/calcChain.xml').
    2. Menstandarkan tag formula kosong pada worksheet xml: mengubah '<v />' atau '<v></v>'
       menjadi '<v>0</v>'. OpenXML strict validator pada Microsoft Excel menolak '<v />' kosong
       di dalam tag formula yang sering memicu error repair dan penghapusan border/rumus.

    Args:
        file_path: Path absolut atau relatif ke berkas Excel yang baru disimpan.
    """
    if not file_path or not os.path.isfile(file_path):
        return

    import tempfile
    import zipfile

    temp_path = file_path + ".tmp_clean"
    try:
        with zipfile.ZipFile(file_path, "r") as zin, zipfile.ZipFile(
            temp_path, "w", compression=zipfile.ZIP_DEFLATED
        ) as zout:
            for item in zin.infolist():
                # Lewati calcChain.xml agar Excel membuat calculation chain baru secara dinamis
                if item.filename == "xl/calcChain.xml":
                    continue

                content = zin.read(item.filename)

                # Hapus relasi ke calcChain di workbook.xml.rels
                if item.filename == "xl/_rels/workbook.xml.rels":
                    text = content.decode("utf-8")
                    text = re.sub(
                        r'<Relationship[^>]*Target=[\"\']calcChain\.xml[\"\'][^>]*/>',
                        "",
                        text,
                    )
                    content = text.encode("utf-8")

                # Hapus override calcChain di [Content_Types].xml
                elif item.filename == "[Content_Types].xml":
                    text = content.decode("utf-8")
                    text = re.sub(
                        r'<Override[^>]*PartName=[\"\']/xl/calcChain\.xml[\"\'][^>]*/>',
                        "",
                        text,
                    )
                    content = text.encode("utf-8")

                # Pastikan atribut applyBorder, applyFill, applyFont, applyNumberFormat aktif pada cellXfs di styles.xml
                # agar Microsoft Excel tidak mengabaikan border sel, warna latar belakang (fill), dan format font
                elif item.filename == "xl/styles.xml":
                    text = content.decode("utf-8")

                    def patch_xf(m: re.Match) -> str:
                        tag = m.group(0)
                        b_m = re.search(r'\bborderId=[\"\'](\d+)[\"\']', tag)
                        if b_m and int(b_m.group(1)) > 0 and 'applyBorder=' not in tag:
                            tag = tag.replace('<xf ', '<xf applyBorder="1" ')
                        f_m = re.search(r'\bfillId=[\"\'](\d+)[\"\']', tag)
                        if f_m and int(f_m.group(1)) > 0 and int(f_m.group(1)) != 1 and 'applyFill=' not in tag:
                            tag = tag.replace('<xf ', '<xf applyFill="1" ')
                        font_m = re.search(r'\bfontId=[\"\'](\d+)[\"\']', tag)
                        if font_m and int(font_m.group(1)) > 0 and 'applyFont=' not in tag:
                            tag = tag.replace('<xf ', '<xf applyFont="1" ')
                        num_m = re.search(r'\bnumFmtId=[\"\'](\d+)[\"\']', tag)
                        if num_m and int(num_m.group(1)) > 0 and 'applyNumberFormat=' not in tag:
                            tag = tag.replace('<xf ', '<xf applyNumberFormat="1" ')
                        return tag

                    def patch_cell_xfs(m: re.Match) -> str:
                        return re.sub(r'<xf\b[^>]*>', patch_xf, m.group(0))

                    text = re.sub(r'<cellXfs\b[^>]*>.*?</cellXfs>', patch_cell_xfs, text, flags=re.DOTALL)
                    content = text.encode("utf-8")

                # Standarkan tag nilai formula kosong pada worksheet
                elif item.filename.startswith("xl/worksheets/sheet") and item.filename.endswith(".xml"):
                    text = content.decode("utf-8")
                    text = re.sub(r'<v\s*/>', '<v>0</v>', text)
                    text = re.sub(r'<v>\s*</v>', '<v>0</v>', text)
                    content = text.encode("utf-8")

                zout.writestr(item, content)

        os.replace(temp_path, file_path)
    except Exception:
        # Jika terjadi kendala saat pasca-proses, pertahankan file asli dan bersihkan temp
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except OSError:
                pass


def preserve_merged_cell_borders(worksheet: Worksheet) -> None:
    """
    Memastikan border perimeter pada merged cell ranges tetap utuh dan terdefinisi
    pada seluruh sel tepi rentang (top, bottom, left, right) agar tidak hilang
    saat diserialisasi oleh openpyxl.
    """
    from openpyxl.styles import Border, Side

    for rng in list(worksheet.merged_cells.ranges):
        min_col, min_row, max_col, max_row = rng.bounds
        top_left = worksheet.cell(min_row, min_col)
        b = top_left.border
        if not b or not any([b.left.style, b.right.style, b.top.style, b.bottom.style]):
            continue

        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cell = worksheet.cell(r, c)
                cb = cell.border
                l_side = b.left if (c == min_col and b.left and b.left.style) else (cb.left if cb and cb.left else Side())
                r_side = b.right if (c == max_col and b.right and b.right.style) else (cb.right if cb and cb.right else Side())
                t_side = b.top if (r == min_row and b.top and b.top.style) else (cb.top if cb and cb.top else Side())
                b_side = b.bottom if (r == max_row and b.bottom and b.bottom.style) else (cb.bottom if cb and cb.bottom else Side())
                cell.border = Border(left=l_side, right=r_side, top=t_side, bottom=b_side)



def format_output_filename(
    pattern: str,
    project_data: Optional[Any] = None,
    template: Optional[Any] = None,
    text_assignments: Optional[Dict[str, str]] = None,
    custom_vars: Optional[Dict[str, str]] = None,
) -> str:
    """
    Menghasilkan nama file dokumen output laporan Excel (.xlsx) berdasarkan pola format
    yang ditentukan di template dengan dukungan substitusi variabel dinamis.

    Variabel standar yang didukung:
    - {site_name} / {site}: Nama/ID site proyek (dibersihkan dari karakter ilegal).
    - {template_name}: Nama konfigurasi template.
    - {date}: Tanggal hari ini format YYYYMMDD (contoh: 20261018).
    - {datetime}: Tanggal & jam format YYYYMMDD_HHMMSS.
    - {doc_name}: Nama sub-dokumen aktif jika ada.
    - {nama_field} / {field_id}: Nilai dari field teks template (contoh: {Nama Site}, {Nomor SPK}).

    Args:
        pattern: String format pola (contoh: "{template_name}_{site_name}_{Nama Site}_{date}").
        project_data: Objek ProjectData jika ada.
        template: Objek TemplateConfig jika ada.
        text_assignments: Dict pemetaan field_id -> teks yang diinputkan pengguna.
        custom_vars: Dict variabel tambahan jika ada.

    Returns:
        str: Nama file berekstensi .xlsx yang valid dan aman untuk sistem berkas OS.
    """
    now = datetime.now()
    date_str = now.strftime("%Y%m%d")
    datetime_str = now.strftime("%Y%m%d_%H%M%S")

    # Siapkan variabel standar
    site_raw = project_data.site_name.strip() if project_data and getattr(project_data, "site_name", None) else ""
    site_slug = re.sub(r'[\\/*?:"<>|\s]+', "_", site_raw) if site_raw else "BTS_Site"

    template_raw = template.name.strip() if template and getattr(template, "name", None) else ""
    template_slug = re.sub(r'[\\/*?:"<>|\s]+', "_", template_raw) if template_raw else "Laporan"

    doc_raw = ""
    if project_data and hasattr(project_data, "get_active_document"):
        active_doc = project_data.get_active_document()
        if active_doc and getattr(active_doc, "name", ""):
            doc_raw = active_doc.name.strip()
    doc_slug = re.sub(r'[\\/*?:"<>|\s]+', "_", doc_raw) if doc_raw else ""

    # Kumpulkan nilai variabel ke dalam lookup dictionary (case-insensitive)
    var_map: Dict[str, str] = {
        "site_name": site_slug,
        "site": site_slug,
        "template_name": template_slug,
        "template": template_slug,
        "date": date_str,
        "datetime": datetime_str,
        "doc_name": doc_slug,
        "doc": doc_slug,
    }

    # Variabel dari field teks template
    combined_texts: Dict[str, str] = {}
    if project_data and getattr(project_data, "text_assignments", None):
        combined_texts.update(project_data.text_assignments)
    if text_assignments:
        combined_texts.update(text_assignments)

    if template and getattr(template, "fields", None):
        for fld in template.fields:
            val = combined_texts.get(fld.field_id, "")
            if not val and getattr(fld, "default_text", ""):
                val = fld.default_text
            clean_val = re.sub(r'[\\/*?:"<>|\s]+', "_", str(val).strip())
            # Simpan berdasarkan field_id dan label
            var_map[fld.field_id.lower()] = clean_val
            var_map[fld.label.strip().lower()] = clean_val
            normalized_label = re.sub(r"\s+", "_", fld.label.strip().lower())
            var_map[normalized_label] = clean_val

    # Tambahkan custom_vars jika ada
    if custom_vars:
        for k, v in custom_vars.items():
            clean_v = re.sub(r'[\\/*?:"<>|\s]+', "_", str(v).strip())
            var_map[k.lower()] = clean_v

    raw_pattern = pattern.strip() if pattern else ""
    if not raw_pattern:
        # Pola fallback standar
        return f"Laporan_{site_slug}_{datetime_str}.xlsx"

    # Substitusi placeholder {placeholder}
    def _replace_match(match: re.Match) -> str:
        key = match.group(1).strip().lower()
        return var_map.get(key, "")

    formatted_name = re.sub(r"\{([^}]+)\}", _replace_match, raw_pattern)

    # Bersihkan karakter terlarang sistem berkas
    clean_name = re.sub(r'[\\/*?:"<>|]+', "_", formatted_name)
    # Hapus garis bawah berulang dan spasi di ujung
    clean_name = re.sub(r"_+", "_", clean_name).strip("_. ")

    # Jika nama kosong setelah substitusi, gunakan default fallback
    if not clean_name:
        clean_name = f"Laporan_{site_slug}_{datetime_str}"

    if not clean_name.lower().endswith(".xlsx"):
        clean_name += ".xlsx"

    return clean_name


